"""AP-0085: el sanitizador elimina metadatos de PDF y XLSX antes de exponerlos."""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter

from src.infrastructure.files.sanitizador_metadatos_archivo import (
    SanitizadorMetadatosArchivo,
)


def _pdf_con_metadatos() -> bytes:
    writer: PdfWriter = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_metadata({"/Author": "Juan Perez", "/Title": "Cedula 123"})
    buf: io.BytesIO = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _xlsx_con_metadatos() -> bytes:
    libro: Workbook = Workbook()
    libro.active["A1"] = "dato"
    libro.properties.creator = "Persona Secreta"
    libro.properties.title = "Reporte interno"
    buf: io.BytesIO = io.BytesIO()
    libro.save(buf)
    return buf.getvalue()


def test_pdf_pierde_metadatos() -> None:
    san: SanitizadorMetadatosArchivo = SanitizadorMetadatosArchivo()
    limpio: bytes = san.sanitizar(_pdf_con_metadatos(), "4123.pdf")
    reader: PdfReader = PdfReader(io.BytesIO(limpio))
    md = reader.metadata
    assert md is None or (md.author is None and md.title is None)


def test_pdf_conserva_paginas() -> None:
    san: SanitizadorMetadatosArchivo = SanitizadorMetadatosArchivo()
    limpio: bytes = san.sanitizar(_pdf_con_metadatos(), "4123.pdf")
    reader: PdfReader = PdfReader(io.BytesIO(limpio))
    assert len(reader.pages) == 1


def test_xlsx_pierde_metadatos() -> None:
    san: SanitizadorMetadatosArchivo = SanitizadorMetadatosArchivo()
    limpio: bytes = san.sanitizar(_xlsx_con_metadatos(), "reporte.xlsx")
    libro: Workbook = load_workbook(io.BytesIO(limpio))
    assert libro.properties.creator in (None, "")
    assert libro.properties.title in (None, "")


def test_tipo_no_soportado_pasa_sin_cambios() -> None:
    san: SanitizadorMetadatosArchivo = SanitizadorMetadatosArchivo()
    original: bytes = b"contenido arbitrario"
    assert san.sanitizar(original, "notas.txt") == original


def test_pdf_corrupto_devuelve_original() -> None:
    san: SanitizadorMetadatosArchivo = SanitizadorMetadatosArchivo()
    basura: bytes = b"esto no es un pdf"
    assert san.sanitizar(basura, "x.pdf") == basura


def test_reporte_excel_sin_metadatos_de_autor() -> None:
    from src.infrastructure.reporting.excel_report_generator import _build_workbook
    data: bytes = _build_workbook("Hoja", ["A", "B"], [["1", "2"]])
    libro: Workbook = load_workbook(io.BytesIO(data))
    assert libro.properties.creator in (None, "")
