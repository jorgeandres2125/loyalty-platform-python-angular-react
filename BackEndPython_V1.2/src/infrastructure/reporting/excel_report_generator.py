from __future__ import annotations

import io
from collections.abc import Mapping, Sequence
from datetime import UTC, date, datetime, time
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.packaging.core import DocumentProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from src.domain.ports.outbound.cifrador_campos import CifradorCampos
from src.infrastructure.reporting.reporte_query import ReporteQuery
from src.shared.constants.cifrado import TABLA_PERFIL_CONTACTO
from src.shared.utils.cifrado_aad import construir_aad

_HEADER_FILL: Final[PatternFill] = PatternFill(
    start_color="1F4E79", end_color="1F4E79", fill_type="solid"
)
_HEADER_FONT: Final[Font] = Font(bold=True, color="FFFFFF")
_HEADER_ALIGN: Final[Alignment] = Alignment(horizontal="center", vertical="center")

_HEADERS_HOJA_VIDA: Final[tuple[str, ...]] = (
    "ARL", "NUMEROIDENTIFIC", "TIPOIDENTIFICAC", "NOMBRERESUMIDO", "FECHANACIMIENTO",
    "COD PAISNACIMIENTO", "PAIS",
    "COD DEPARTNACIMIENT", "DEPARTAMENTO",
    "COD CIUDADNACIMIENT", "CIUDAD",
    "SEXO", "DIRECCIÓN",
    "TELEFONO1", "CELULAR", "EMAIL", "FIRMACONTRATO", "ESTADO", "CANAL",
    "OFICINA-CONCESIONARIO", "INCENTIVOS", "COMISION",
)
_PAIS_DEFAULT: Final[str] = "COLOMBIA"

_HEADERS_INFO_LABORAL: Final[tuple[str, ...]] = (
    "TIPO DOCUMENTO", "EMPLEADO", "ARP", "NIT TERCEROPENSION", "NIT TERCEROSALUD",
    "COD BANCO", "NUMERCUENTBANCO", "TIPOCUENTA",
)

_HEADERS_PLANTILLA: Final[tuple[str, ...]] = (
    "cedula", "correo electrónico", "direccion", "telefono", "celular",
    "sexo", "fechaNacimiento (dd/mm/aaaa)", "Activo(si/no)",
    "Ciudad", "Departamento", "Direccion Envío", "Ciudad Envío",
    "Departamento envío", "Teléfono envío", "Celular envío", "Nombre contacto",
)

_HEADERS_TRIBUTARIA: Final[tuple[str, ...]] = (
    "Tipo de Documento", "Cedula", "Salarios y prestaciones sociales por vinculación", "Comisiones",
    "Honorarios", "Arrendamientos", "Ingresos por pensión, invalidez vejez o muerte",
    "Actividades deportivas y otras actividades", "Agropecuario, silvicultura y pesca",
    "Comercio al por mayor", "Comercio al por menor",
    "Comercio de vehículos automotores, accesorios y productos conexos",
    "Construcción", "Electricidad, gas y vapor", "Fabricación de productos minerales y otros",
    "Fabricación de sustancias químicas", "Industria de la madera, corcho y papel",
    "Manufactura de alimentos", "Manufactura textiles, prendas de vestir y cuero",
    "Minería", "Servicio de transporte, almacenamiento y comunicaciones",
    "Servicios de Hoteles, restaurantes y similares", "Servicios financieros",
    "Dividendos", "Rendimientos financieros", "otros", "Otros cual",
    "Debes presentar declaración de renta", "Mis ingresos totales superan",
    "EPS", "Fondo de pensiones", "ARL", "TOTAL TRIBUTARIO",
)

# Columnas 1-based de _HEADERS_TRIBUTARIA con valores decimales (cols 3-26 + 33)
_TRIBUTARIA_NUMERIC_COLS: Final[frozenset[int]] = frozenset(range(3, 27)) | frozenset({33})

_HEADERS_MIGRADOS: Final[tuple[str, ...]] = (
    "cmid", "L111NID", "L111TID", "L111NOM", "L111NCA", "L111DCA", "L111ALI",
    "L111PCP", "L111CEL", "L111TE2", "L111DIR", "L111CIU", "L111DEP",
    "L111FNA", "L111CTA", "L111TCT", "L111CBC", "L111DBC", "L111INI",
    "status", "created", "changed",
)

_HEADERS_DEFAULT: Final[tuple[str, ...]] = ("Nombre", "Email", "Último acceso")

_HEADERS_ESTADOS: Final[tuple[str, ...]] = (
    "CC", "ESTADO", "CAMPOS FALTANTES", "DOCUMENTOS FALTANTES", "FECHA DE COMPLETADO",
)


def _date_to_unix(d: date | None, *, end_of_day: bool = False) -> int | None:
    if d is None:
        return None
    tiempo: time = time(23, 59, 59) if end_of_day else time(0, 0, 0)
    return int(datetime.combine(d, tiempo, tzinfo=UTC).timestamp())


def _decimal2(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def _to_jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _limpiar_metadatos_workbook(wb: Workbook) -> None:
    """AP-0085: elimina metadatos del libro (autor, fechas) antes de exponerlo."""
    props: DocumentProperties = DocumentProperties()
    props.creator = ""
    props.lastModifiedBy = ""
    props.title = None
    props.subject = None
    props.description = None
    props.keywords = None
    props.category = None
    wb.properties = props


def _build_workbook(
    sheet_name: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    numeric_cols: frozenset[int] | None = None,
) -> bytes:
    wb: Workbook = Workbook()
    ws: Any = wb.active
    ws.title = sheet_name
    for col_idx, header in enumerate(headers, start=1):
        cell: Any = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 2)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            data_cell: Any = ws.cell(row=r_idx, column=c_idx, value=value)
            if numeric_cols and c_idx in numeric_cols:
                data_cell.number_format = "0.00"
    ws.freeze_panes = "A2"
    _limpiar_metadatos_workbook(wb)
    buf: io.BytesIO = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bool_si_no(value: Any) -> str:
    if value in (1, True, "1", "SI", "Si", "si"):
        return "SI"
    return "NO"


def _genero_letter(value: Any) -> str:
    if value in (1, "1", "M", "m"):
        return "M"
    if value in (0, "0", "F", "f"):
        return "F"
    return ""


def _genero_label(value: Any) -> str:
    if value in (1, "1", "M", "m"):
        return "Masculino"
    if value in (0, "0", "F", "f"):
        return "Femenino"
    return ""


def _cuenta_label(value: Any) -> str:
    if value in (0, "0", "Ahorros", "ahorros"):
        return "Ahorros"
    if value in (1, "1", "Corriente", "corriente"):
        return "Corriente"
    return ""


def _estado_hoja_vida(estado_contacto: Any, estado_tributario: Any) -> str:
    if estado_contacto is not None and estado_tributario is not None:
        if int(estado_contacto) == 1 and int(estado_tributario) == 1:
            return "SI"
        if int(estado_tributario) == 2:
            return "PENDIENTE DOC"
    if estado_contacto is not None and int(estado_contacto) == 1 and not estado_tributario:
        return "SI"
    return "NO"


class ExcelReportGenerator:
    def __init__(self, session: AsyncSession, cipher: CifradorCampos) -> None:
        self._session: AsyncSession = session
        self._cipher: CifradorCampos = cipher

    def _descifrar_contacto(self, valor_enc: Any, columna: str, pk: Any) -> str | None:
        """Descifra una columna restringida de users_perfil_contacto leída como
        VARBINARY por SQL crudo. pk = numero_documento (para el AAD)."""
        if valor_enc is None or pk is None:
            return None
        return self._cipher.descifrar(
            valor_enc, aad=construir_aad(TABLA_PERFIL_CONTACTO, columna, str(pk))
        )

    def _date_filters(
        self, fecha_inicio: date | None, fecha_fin: date | None,
    ) -> tuple[int, int]:
        ini: int = _date_to_unix(fecha_inicio) or 0
        fin: int = _date_to_unix(fecha_fin, end_of_day=True) or int(datetime.now(UTC).timestamp())
        return ini, fin

    @staticmethod
    def _subprograma_clause(subprograma: int | None) -> str:
        return (
            " AND uc.comisionista_subprograma_id = :subprograma"
            if subprograma is not None
            else ""
        )

    @staticmethod
    def _add_subprograma(params: dict[str, Any], subprograma: int | None) -> None:
        if subprograma is not None:
            params["subprograma"] = subprograma

    async def _fetch(self, sql: str, params: Mapping[str, Any]) -> list[Mapping[str, Any]]:
        result = await self._session.execute(text(sql), params)
        return [dict(row._mapping) for row in result.all()]

    async def _fetch_scalar(self, sql: str, params: Mapping[str, Any]) -> int:
        result = await self._session.execute(text(sql), params)
        value: Any = result.scalar()
        return int(value or 0)

    # ---------- Builders por reporte ----------
    def _build_hoja_vida(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> ReporteQuery:
        ini: int
        fin: int
        ini, fin = self._date_filters(fecha_inicio, fecha_fin)
        sql_body: str = f"""
        SELECT
            tarl.name        AS arl,
            uc.numero_documento AS cedula,
            uc.tipo_documento,
            uc.nombre_completo,
            CAST(uc.fecha_nacimiento AS date) AS fecha_nacimiento,
            d.pid            AS codigo_pais,
            uc.departamento  AS codigo_departamento,
            d.departamento   AS departamento_nombre,
            uc.ciudad        AS codigo_ciudad,
            c.ciudad         AS ciudad_nombre,
            uc.genero,
            uc.direccion_enc,
            uc.telefono_enc,
            uc.celular_enc,
            us.mail,
            um.FIRMA_CONTRATO,
            uc.estado        AS estado_contacto,
            ut.estado        AS estado_tributario,
            can.id_canales,
            can.nom_canales,
            ofi.id_oficinas,
            ofi.nom_oficinas,
            uc.incentivos,
            uc.requiere_comision,
            us.created       AS _orden
        FROM users us
        INNER JOIN users_perfil_contacto uc ON us.name = uc.numero_documento
        LEFT JOIN users_migracion um ON um.L222NID = uc.numero_documento
        LEFT JOIN canales can ON can.cod_canales = uc.cod_canales
        LEFT JOIN oficinas ofi ON ofi.cod_oficinas = uc.cod_oficinas
        LEFT JOIN users_perfil_tributario ut ON us.name = ut.numero_documento AND ut.deleted = 0
        LEFT JOIN departamentos d ON d.did = TRY_CAST(uc.departamento AS int)
        LEFT JOIN ciudades c      ON c.cid = TRY_CAST(uc.ciudad AS int)
        LEFT JOIN taxonomy_term_data tarl ON tarl.tid = (
            SELECT TOP 1 CASE WHEN ut2.arl IS NULL THEN TRY_CAST(um.ARL AS int) ELSE ut2.arl END
            FROM users_perfil_tributario ut2
            WHERE ut2.numero_documento = uc.numero_documento
        )
        WHERE us.created BETWEEN :ini AND :fin
          AND uc.comisionista_programa_id = :programa{self._subprograma_clause(subprograma)}
        """
        order_by: str = " ORDER BY _orden DESC "
        params: dict[str, Any] = {"ini": ini, "fin": fin, "programa": programa}
        self._add_subprograma(params, subprograma)

        def transformer(row: Mapping[str, Any]) -> list[Any]:
            pk: Any = row.get("cedula")
            return [
                row.get("arl"),
                row.get("cedula"),
                row.get("tipo_documento"),
                row.get("nombre_completo"),
                row.get("fecha_nacimiento"),
                row.get("codigo_pais"),
                _PAIS_DEFAULT if row.get("codigo_pais") is not None else "",
                row.get("codigo_departamento"),
                row.get("departamento_nombre"),
                row.get("codigo_ciudad"),
                row.get("ciudad_nombre"),
                _genero_label(row.get("genero")),
                self._descifrar_contacto(row.get("direccion_enc"), "direccion", pk),
                self._descifrar_contacto(row.get("telefono_enc"), "telefono", pk),
                self._descifrar_contacto(row.get("celular_enc"), "celular", pk),
                row.get("mail"),
                _bool_si_no(row.get("FIRMA_CONTRATO")),
                _estado_hoja_vida(row.get("estado_contacto"), row.get("estado_tributario")),
                f"{row.get('id_canales') or ''}-{row.get('nom_canales') or ''}",
                f"{row.get('id_oficinas') or ''}-{row.get('nom_oficinas') or ''}",
                _bool_si_no(row.get("incentivos")),
                _bool_si_no(row.get("requiere_comision")),
            ]

        return ReporteQuery(
            sheet_name="Hoja de vida",
            headers=_HEADERS_HOJA_VIDA,
            sql_body=sql_body,
            order_by=order_by,
            params=params,
            transformer=transformer,
        )

    def _build_info_laboral(
        self,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> ReporteQuery:
        ini: int
        fin: int
        ini, fin = self._date_filters(fecha_inicio, fecha_fin)
        sql_body: str = """
        SELECT
            uc.tipo_documento,
            us.name                  AS empleado,
            tarl.name                AS arl,
            tafp.description         AS afp,
            teps.description         AS eps,
            banco.description        AS banco,
            uc.numero_documento      AS _pk_doc,
            uc.numero_de_cuenta_enc,
            uc.tipo_de_cuenta_enc,
            us.created               AS _orden
        FROM users us
        INNER JOIN users_perfil_contacto uc ON us.name = uc.numero_documento
        LEFT JOIN users_perfil_tributario ut ON us.name = ut.numero_documento AND ut.deleted = 0
        LEFT JOIN users_migracion um ON um.L222NID = uc.numero_documento
        LEFT JOIN taxonomy_term_data tarl ON tarl.tid = (
            SELECT TOP 1 CASE WHEN ut2.arl IS NULL THEN TRY_CAST(um.ARL AS int) ELSE ut2.arl END
            FROM users_perfil_tributario ut2 WHERE ut2.numero_documento = uc.numero_documento
        )
        LEFT JOIN taxonomy_term_data teps ON teps.tid = (
            SELECT TOP 1 CASE WHEN ut2.eps IS NULL THEN TRY_CAST(um.EPS AS int) ELSE ut2.eps END
            FROM users_perfil_tributario ut2 WHERE ut2.numero_documento = uc.numero_documento
        )
        LEFT JOIN taxonomy_term_data tafp ON tafp.tid = (
            SELECT TOP 1 CASE WHEN ut2.afp IS NULL THEN TRY_CAST(um.AFP AS int) ELSE ut2.afp END
            FROM users_perfil_tributario ut2 WHERE ut2.numero_documento = uc.numero_documento
        )
        LEFT JOIN taxonomy_term_data banco ON banco.tid = uc.banco
        WHERE us.created BETWEEN :ini AND :fin
        """
        order_by: str = " ORDER BY _orden DESC "
        params: dict[str, Any] = {"ini": ini, "fin": fin}

        def transformer(row: Mapping[str, Any]) -> list[Any]:
            pk: Any = row.get("_pk_doc")
            cuenta: str | None = self._descifrar_contacto(
                row.get("numero_de_cuenta_enc"), "numero_de_cuenta", pk
            )
            tipo: str | None = self._descifrar_contacto(
                row.get("tipo_de_cuenta_enc"), "tipo_de_cuenta", pk
            )
            return [
                row.get("tipo_documento"),
                row.get("empleado"),
                row.get("arl"),
                row.get("afp"),
                row.get("eps"),
                row.get("banco"),
                cuenta,
                _cuenta_label(tipo),
            ]

        return ReporteQuery(
            sheet_name="Info laboral",
            headers=_HEADERS_INFO_LABORAL,
            sql_body=sql_body,
            order_by=order_by,
            params=params,
            transformer=transformer,
        )

    def _build_plantilla_participantes(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> ReporteQuery:
        ini: int
        fin: int
        ini, fin = self._date_filters(fecha_inicio, fecha_fin)
        sql_body: str = f"""
        SELECT
            uc.numero_documento,
            us.mail,
            uc.direccion_enc,
            uc.telefono_enc,
            uc.celular_enc,
            uc.genero,
            CAST(uc.fecha_nacimiento AS date) AS fecha_nacimiento,
            c.ciudad        AS ciudad_nombre,
            d.departamento  AS departamento_nombre,
            uc.nombre_completo,
            us.created      AS _orden
        FROM users us
        INNER JOIN users_perfil_contacto uc ON us.name = uc.numero_documento
        LEFT JOIN ciudades c       ON c.cid = TRY_CAST(uc.ciudad AS int)
        LEFT JOIN departamentos d  ON d.did = TRY_CAST(uc.departamento AS int)
        WHERE us.created BETWEEN :ini AND :fin
          AND uc.comisionista_programa_id = :programa{self._subprograma_clause(subprograma)}
        """
        order_by: str = " ORDER BY _orden DESC "
        params: dict[str, Any] = {"ini": ini, "fin": fin, "programa": programa}
        self._add_subprograma(params, subprograma)

        def transformer(row: Mapping[str, Any]) -> list[Any]:
            pk: Any = row.get("numero_documento")
            direccion: str | None = self._descifrar_contacto(
                row.get("direccion_enc"), "direccion", pk
            )
            telefono: str | None = self._descifrar_contacto(row.get("telefono_enc"), "telefono", pk)
            celular: str | None = self._descifrar_contacto(row.get("celular_enc"), "celular", pk)
            return [
                row.get("numero_documento"),
                row.get("mail"),
                direccion,
                telefono,
                celular,
                _genero_label(row.get("genero")),
                row.get("fecha_nacimiento"),
                "N/A",
                row.get("ciudad_nombre"),
                row.get("departamento_nombre"),
                direccion,
                row.get("ciudad_nombre"),
                row.get("departamento_nombre"),
                telefono,
                celular,
                row.get("nombre_completo"),
            ]

        return ReporteQuery(
            sheet_name="Plantilla participantes",
            headers=_HEADERS_PLANTILLA,
            sql_body=sql_body,
            order_by=order_by,
            params=params,
            transformer=transformer,
        )

    def _build_info_tributaria(
        self,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> ReporteQuery:
        ini: int
        fin: int
        ini, fin = self._date_filters(fecha_inicio, fecha_fin)
        sql_body: str = """
        SELECT
            uc.tipo_documento,
            us.name AS cedula,
            ut.salarios_y_prestaciones, ut.comisiones, ut.honorarios,
            ut.arrendamientos, ut.ingresos_por_pension, ut.actividades_deportivas,
            ut.agropecuario, ut.comercio_mayor, ut.comercio_menor, ut.comercio_vehiculos,
            ut.construccion, ut.electricidad_gas, ut.fabricacion_minerales,
            ut.fabricacion_sustancias, ut.industria_madera, ut.manufactura_alimentos,
            ut.manufactura_textiles, ut.mineria, ut.servicios_transporte,
            ut.servicios_de_hoteles, ut.servicios_financieros, ut.renta_dividendos,
            ut.renta_financiero, ut.otro, ut.otros_cual, ut.declaracion_renta,
            ut.ingresos_mayor, ut.total,
            teps.description AS eps_code,
            tafp.description AS afp_code,
            tarl.description AS arl_code,
            us.created       AS _orden
        FROM users us
        INNER JOIN users_perfil_contacto uc ON us.name = uc.numero_documento
        LEFT JOIN users_perfil_tributario ut ON us.name = ut.numero_documento AND ut.deleted = 0
        LEFT JOIN users_migracion um ON um.L222NID = uc.numero_documento
        LEFT JOIN taxonomy_term_data tarl ON tarl.tid = (
            SELECT TOP 1 CASE WHEN ut2.arl IS NULL THEN TRY_CAST(um.ARL AS int) ELSE ut2.arl END
            FROM users_perfil_tributario ut2 WHERE ut2.numero_documento = uc.numero_documento
        )
        LEFT JOIN taxonomy_term_data teps ON teps.tid = (
            SELECT TOP 1 CASE WHEN ut2.eps IS NULL THEN TRY_CAST(um.EPS AS int) ELSE ut2.eps END
            FROM users_perfil_tributario ut2 WHERE ut2.numero_documento = uc.numero_documento
        )
        LEFT JOIN taxonomy_term_data tafp ON tafp.tid = (
            SELECT TOP 1 CASE WHEN ut2.afp IS NULL THEN TRY_CAST(um.AFP AS int) ELSE ut2.afp END
            FROM users_perfil_tributario ut2 WHERE ut2.numero_documento = uc.numero_documento
        )
        WHERE us.created BETWEEN :ini AND :fin
        """
        order_by: str = " ORDER BY _orden DESC "
        params: dict[str, Any] = {"ini": ini, "fin": fin}

        def transformer(row: Mapping[str, Any]) -> list[Any]:
            return [
                row.get("tipo_documento"),
                row.get("cedula"),
                _decimal2(row.get("salarios_y_prestaciones")),
                _decimal2(row.get("comisiones")),
                _decimal2(row.get("honorarios")),
                _decimal2(row.get("arrendamientos")),
                _decimal2(row.get("ingresos_por_pension")),
                _decimal2(row.get("actividades_deportivas")),
                _decimal2(row.get("agropecuario")),
                _decimal2(row.get("comercio_mayor")),
                _decimal2(row.get("comercio_menor")),
                _decimal2(row.get("comercio_vehiculos")),
                _decimal2(row.get("construccion")),
                _decimal2(row.get("electricidad_gas")),
                _decimal2(row.get("fabricacion_minerales")),
                _decimal2(row.get("fabricacion_sustancias")),
                _decimal2(row.get("industria_madera")),
                _decimal2(row.get("manufactura_alimentos")),
                _decimal2(row.get("manufactura_textiles")),
                _decimal2(row.get("mineria")),
                _decimal2(row.get("servicios_transporte")),
                _decimal2(row.get("servicios_de_hoteles")),
                _decimal2(row.get("servicios_financieros")),
                _decimal2(row.get("renta_dividendos")),
                _decimal2(row.get("renta_financiero")),
                _decimal2(row.get("otro")),
                row.get("otros_cual"),
                "Si" if row.get("declaracion_renta") else "No",
                "Si" if row.get("ingresos_mayor") else "No",
                row.get("eps_code"),
                row.get("afp_code"),
                row.get("arl_code"),
                _decimal2(row.get("total")),
            ]

        return ReporteQuery(
            sheet_name="Info tributaria",
            headers=_HEADERS_TRIBUTARIA,
            sql_body=sql_body,
            order_by=order_by,
            params=params,
            transformer=transformer,
            numeric_cols=_TRIBUTARIA_NUMERIC_COLS,
        )

    def _build_usuarios_migrados(
        self,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> ReporteQuery:
        ini: int
        fin: int
        ini, fin = self._date_filters(fecha_inicio, fecha_fin)
        sql_body: str = """
        SELECT
            cmid, L111NID, L111TID, L111NOM, L111NCA, L111DCA, L111ALI, L111PCP,
            L111CEL, L111TE2, L111DIR, L111CIU, L111DEP, L111FNA, L111CTA,
            L111TCT, L111CBC, L111DBC, L111INI, status, created, changed
        FROM comisionistas_sufi_migrado
        WHERE created BETWEEN :ini AND :fin
        """
        order_by: str = " ORDER BY created DESC "
        params: dict[str, Any] = {"ini": ini, "fin": fin}

        def transformer(row: Mapping[str, Any]) -> list[Any]:
            return [row.get(header) for header in _HEADERS_MIGRADOS]

        return ReporteQuery(
            sheet_name="Usuarios migrados",
            headers=_HEADERS_MIGRADOS,
            sql_body=sql_body,
            order_by=order_by,
            params=params,
            transformer=transformer,
        )

    def _build_default(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> ReporteQuery:
        ini: int
        fin: int
        ini, fin = self._date_filters(fecha_inicio, fecha_fin)
        sql_body: str = f"""
        SELECT us.name, us.mail, us.access, us.created AS _orden
        FROM users us
        INNER JOIN users_perfil_contacto uc ON us.name = uc.numero_documento
        LEFT JOIN users_perfil_tributario ut ON us.name = ut.numero_documento AND ut.deleted = 0
        WHERE us.created BETWEEN :ini AND :fin
          AND uc.comisionista_programa_id = :programa{self._subprograma_clause(subprograma)}
        """
        order_by: str = " ORDER BY _orden DESC "
        params: dict[str, Any] = {"ini": ini, "fin": fin, "programa": programa}
        self._add_subprograma(params, subprograma)

        def transformer(row: Mapping[str, Any]) -> list[Any]:
            access: Any = row.get("access")
            ultimo: str = ""
            if access:
                try:
                    ultimo = datetime.fromtimestamp(int(access), tz=UTC).strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                except (ValueError, OSError):
                    ultimo = ""
            return [row.get("name"), row.get("mail"), ultimo]

        return ReporteQuery(
            sheet_name="Listado",
            headers=_HEADERS_DEFAULT,
            sql_body=sql_body,
            order_by=order_by,
            params=params,
            transformer=transformer,
        )

    def _build_estados_perfiles(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
        cedula: str | None,
        estado: int | None,
    ) -> ReporteQuery:
        ini: int
        fin: int
        ini, fin = self._date_filters(fecha_inicio, fecha_fin)
        clauses: list[str] = [
            "u.created BETWEEN :ini AND :fin",
            "uc.comisionista_programa_id = :programa",
        ]
        params: dict[str, Any] = {"ini": ini, "fin": fin, "programa": programa}
        if subprograma is not None:
            clauses.append("uc.comisionista_subprograma_id = :subprograma")
            params["subprograma"] = subprograma
        if estado is not None and estado != 3:
            clauses.append("(ut.estado = :est OR uc.estado = :est)")
            params["est"] = estado
        if cedula:
            clauses.append("u.name = :cedula")
            params["cedula"] = cedula
        where: str = " AND ".join(clauses)
        sql_body: str = f"""
        SELECT
            u.name AS cedula,
            uc.estado          AS estado_contacto,
            ut.estado          AS estado_financiero,
            uc.fecha_completado,
            uc.tipo_documento, uc.nombre_completo, uc.genero, uc.fecha_nacimiento,
            uc.telefono_enc, uc.celular_enc, uc.direccion_enc, uc.departamento, uc.ciudad,
            uc.concesionario, uc.tipo_de_cuenta_enc, uc.banco, uc.numero_de_cuenta_enc,
            uc.acepto_habeas_data,
            ue.acepto_terminos_y_condiciones,
            ut.declaracion_renta, ut.ingresos_mayor, ut.eps, ut.afp, ut.arl,
            u.created AS _orden
        FROM users u
        INNER JOIN users_perfil_contacto uc  ON u.name = uc.numero_documento
        INNER JOIN users_perfil_emocional ue ON u.name = ue.numero_documento
        LEFT  JOIN users_perfil_tributario ut ON u.name = ut.numero_documento AND ut.deleted = 0
        WHERE {where}
        """
        order_by: str = " ORDER BY _orden DESC "

        def transformer(row: Mapping[str, Any]) -> list[Any]:
            pk: Any = row.get("cedula")
            fila: dict[str, Any] = dict(row)
            fila["numero_de_cuenta"] = self._descifrar_contacto(
                row.get("numero_de_cuenta_enc"), "numero_de_cuenta", pk
            )
            fila["tipo_de_cuenta"] = self._descifrar_contacto(
                row.get("tipo_de_cuenta_enc"), "tipo_de_cuenta", pk
            )
            for _col in ("telefono", "celular", "direccion"):
                fila[_col] = self._descifrar_contacto(row.get(f"{_col}_enc"), _col, pk)
            faltantes: list[str] = ExcelReportGenerator._campos_faltantes(fila)
            estado_label: str
            if not faltantes:
                estado_label = "Completo"
            elif len(faltantes) == 1 and faltantes[0].startswith("Documentación"):
                estado_label = "Pendiente por Documentación"
            else:
                estado_label = "Incompleto"
            return [
                row.get("cedula"),
                estado_label,
                ", ".join(faltantes),
                "N/A",
                row.get("fecha_completado") if estado_label != "Incompleto" else "",
            ]

        return ReporteQuery(
            sheet_name="Estados de perfiles",
            headers=_HEADERS_ESTADOS,
            sql_body=sql_body,
            order_by=order_by,
            params=params,
            transformer=transformer,
        )

    # ---------- Excel (sin paginar) ----------
    async def _to_excel(self, query: ReporteQuery) -> bytes:
        sql: str = query.sql_body + query.order_by
        rows_raw: list[Mapping[str, Any]] = await self._fetch(sql, query.params)
        rows: list[list[Any]] = [query.transformer(row) for row in rows_raw]
        return _build_workbook(query.sheet_name, query.headers, rows, query.numeric_cols)

    # ---------- Preview (paginado) ----------
    async def _to_preview(
        self, query: ReporteQuery, page: int, page_size: int,
    ) -> tuple[list[str], list[list[Any]], int]:
        count_sql: str = f"SELECT COUNT(*) AS total FROM ({query.sql_body}) X"
        total: int = await self._fetch_scalar(count_sql, query.params)
        offset: int = (page - 1) * page_size
        page_sql: str = (
            f"{query.sql_body}{query.order_by} OFFSET :_off ROWS FETCH NEXT :_ps ROWS ONLY"
        )
        params_page: dict[str, Any] = {**query.params, "_off": offset, "_ps": page_size}
        rows_raw: list[Mapping[str, Any]] = await self._fetch(page_sql, params_page)
        rows: list[list[Any]] = [
            [_to_jsonable(valor) for valor in query.transformer(row)] for row in rows_raw
        ]
        return list(query.headers), rows, total

    # ---------- Reporte 0: Hoja de vida ----------
    async def generar_hoja_vida_async(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> bytes:
        return await self._to_excel(
            self._build_hoja_vida(programa, subprograma, fecha_inicio, fecha_fin)
        )

    # ---------- Reporte 1: Información Laboral ----------
    async def generar_info_laboral_async(
        self,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> bytes:
        return await self._to_excel(self._build_info_laboral(fecha_inicio, fecha_fin))

    # ---------- Reporte 2: Plantilla participantes ----------
    async def generar_plantilla_participantes_async(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> bytes:
        return await self._to_excel(
            self._build_plantilla_participantes(programa, subprograma, fecha_inicio, fecha_fin)
        )

    # ---------- Reporte 3: Información tributaria ----------
    async def generar_info_tributaria_async(
        self,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> bytes:
        return await self._to_excel(self._build_info_tributaria(fecha_inicio, fecha_fin))

    # ---------- Reporte 4: Usuarios migrados ----------
    async def generar_usuarios_migrados_async(
        self,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> bytes:
        return await self._to_excel(self._build_usuarios_migrados(fecha_inicio, fecha_fin))

    # ---------- Reporte 5: Default (listado básico) ----------
    async def generar_default_async(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
    ) -> bytes:
        return await self._to_excel(
            self._build_default(programa, subprograma, fecha_inicio, fecha_fin)
        )

    # ---------- Reporte 6: Estados de perfiles ----------
    async def generar_estados_perfiles_async(
        self,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
        cedula: str | None,
        estado: int | None,
    ) -> bytes:
        return await self._to_excel(
            self._build_estados_perfiles(
                programa, subprograma, fecha_inicio, fecha_fin, cedula, estado
            )
        )

    # ---------- Preview unificado ----------
    async def obtener_preview_async(
        self,
        tipo: int,
        programa: int,
        subprograma: int | None,
        fecha_inicio: date | None,
        fecha_fin: date | None,
        cedula: str | None,
        estado: int | None,
        page: int,
        page_size: int,
    ) -> tuple[list[str], list[list[Any]], int]:
        query: ReporteQuery
        if tipo == 0:
            query = self._build_hoja_vida(programa, subprograma, fecha_inicio, fecha_fin)
        elif tipo == 1:
            query = self._build_info_laboral(fecha_inicio, fecha_fin)
        elif tipo == 2:
            query = self._build_plantilla_participantes(
                programa, subprograma, fecha_inicio, fecha_fin
            )
        elif tipo == 3:
            query = self._build_info_tributaria(fecha_inicio, fecha_fin)
        elif tipo == 4:
            query = self._build_usuarios_migrados(fecha_inicio, fecha_fin)
        elif tipo == 6:
            query = self._build_estados_perfiles(
                programa, subprograma, fecha_inicio, fecha_fin, cedula, estado
            )
        else:
            query = self._build_default(programa, subprograma, fecha_inicio, fecha_fin)
        return await self._to_preview(query, page, page_size)

    @staticmethod
    def _campos_faltantes(row: Mapping[str, Any]) -> list[str]:
        contacto: Mapping[str, str] = {
            "tipo_documento": "Tipo de Documento",
            "nombre_completo": "Nombre Completo",
            "genero": "Género",
            "fecha_nacimiento": "Fecha de nacimiento",
            "telefono": "Teléfono",
            "celular": "Celular",
            "direccion": "Dirección",
            "departamento": "Departamento",
            "ciudad": "Ciudad",
            "concesionario": "Concesionario",
            "tipo_de_cuenta": "Tipo de Cuenta",
            "banco": "Banco",
            "numero_de_cuenta": "Número de Cuenta",
            "acepto_habeas_data": "Habeas Data",
        }
        emocional: Mapping[str, str] = {
            "acepto_terminos_y_condiciones": "Acepto términos y condiciones",
        }
        tributario: Mapping[str, str] = {
            "declaracion_renta": "Declaración de renta",
            "ingresos_mayor": "Ingresos totales",
            "eps": "EPS",
            "afp": "Fondo de pensiones",
            "arl": "ARL",
        }
        faltantes: list[str] = []
        for key, label in contacto.items():
            if row.get(key) in (None, ""):
                faltantes.append(label)
        for key, label in emocional.items():
            if row.get(key) in (None, ""):
                faltantes.append(label)
        if row.get("estado_financiero") is not None:
            for key, label in tributario.items():
                if row.get(key) in (None, ""):
                    faltantes.append(label)
        return faltantes
